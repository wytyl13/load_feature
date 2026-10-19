from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


from workstreams_load_characteristic_model_v2 import data_io
from workstreams_load_characteristic_model_v2.schemas import (
    AccountingRole,
    Official2025Row,
    OfficialMappingEntry,
)


class DataMappingTests(unittest.TestCase):
    @staticmethod
    def _write_excluded_energy_fixture(
        path: Path,
        *,
        omit_label: str | None = None,
        duplicate_label: str | None = None,
        negative_label: str | None = None,
        deduction_exceeds_total: bool = False,
    ) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        # Regression fixture: a bare 2020 appears in metadata before the real
        # ``2020年`` column header and must not be treated as the data column.
        worksheet.cell(1, 1, 2020)
        worksheet.cell(2, 1, "单位：万千瓦时")
        worksheet.cell(3, 1, "指标名称")
        years = tuple(range(2020, 2026))
        for column, year in enumerate(years, start=2):
            worksheet.cell(3, column, f"{year}年")
        labels = [
            "全社会用电总计",
            "其中：电厂生产全部耗用电量",
            "线路损失电量",
            "抽水蓄能抽水耗用电量",
        ]
        labels = [label for label in labels if label != omit_label]
        if duplicate_label is not None:
            labels.append(duplicate_label)
        for row_number, label in enumerate(labels, start=4):
            worksheet.cell(row_number, 1, label)
            for column, year in enumerate(years, start=2):
                offset = year - 2020
                if label == "全社会用电总计":
                    value = 1_000.0 + 100.0 * offset
                elif label == "其中：电厂生产全部耗用电量":
                    value = 50.0 + offset
                elif label == "线路损失电量":
                    value = 20.0 + offset
                else:
                    value = 10.0 + offset
                if deduction_exceeds_total and label == "其中：电厂生产全部耗用电量":
                    value = 2_000.0
                if label == negative_label:
                    value = -1.0
                cell = worksheet.cell(row_number, column, value)
                cell.number_format = "0.0000"
        workbook.save(path)
        workbook.close()

    def test_clean_text_and_sample_sheet_parsing(self) -> None:
        parsed = data_io.parse_sample_sheet_name(" 第一产业-农业-用户1 ")
        self.assertEqual(parsed.sheet_name, "第一产业-农业-用户1")
        self.assertEqual(parsed.broad_sector, "第一产业")
        self.assertEqual(parsed.industry, "农业")
        self.assertEqual(parsed.user_number, 1)

        residential = data_io.parse_sample_sheet_name("居民生活-市中-用户2")
        self.assertEqual(residential.broad_sector, "居民生活")
        self.assertEqual(residential.industry, "居民生活")
        self.assertEqual(residential.user_number, 2)

        with self.assertRaisesRegex(ValueError, "invalid sample sheet name"):
            data_io.parse_sample_sheet_name("第二产业-没有用户编号")

    def test_sha256_manifest_detects_change(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "input.bin"
            source.write_bytes(b"approved-input")
            manifest = data_io.build_input_manifest({"fixture": source})
            self.assertEqual(len(manifest), 1)
            self.assertEqual(len(manifest[0].sha256), 64)
            data_io.verify_input_manifest(manifest)

            source.write_bytes(b"tampered-input")
            with self.assertRaisesRegex(ValueError, "changed"):
                data_io.verify_input_manifest(manifest)

    def test_static_sample_mapping_is_unique_and_complete(self) -> None:
        mapping = data_io.load_sample_mapping()
        self.assertEqual(len(mapping), 25)
        self.assertTrue(mapping["industry"].is_unique)
        self.assertEqual(int(mapping["is_traditional"].sum()), 24)
        self.assertEqual(
            set(mapping.loc[mapping["is_traditional"], "broad_sector"]),
            {"第一产业", "第二产业", "第三产业", "居民生活"},
        )
        charging = mapping.loc[mapping["industry"].eq("充换电服务业")].iloc[0]
        self.assertEqual(charging["modeling_role"], "后续新型负荷模块样本")
        self.assertFalse(bool(charging["is_traditional"]))
        self.assertNotIn("run_group", mapping.columns)
        self.assertNotIn("parameter_group", mapping.columns)

    def test_static_official_mapping_enumerates_all_source_rows(self) -> None:
        entries = data_io._load_official_mapping_entries()  # exact config audit
        self.assertEqual(len(entries), 135)
        self.assertEqual([entry.source_row for entry in entries], list(range(16, 151)))
        self.assertEqual(len({entry.official_label for entry in entries}), 135)
        self.assertTrue(
            {entry.classification for entry in entries}.issuperset(
                {"汇总项", "新型负荷", "排除项"}
            )
        )
        split_labels = {
            entry.official_label
            for entry in entries
            if entry.accounting_role is AccountingRole.SPLIT_PARENT
        }
        self.assertEqual(
            split_labels,
            {
                "1.电力、热力生产和供应业",
                "2.互联网和相关服务",
                "六、批发和零售业",
                "4.教育、文化、体育和娱乐业",
            },
        )

    def test_synthetic_split_parent_and_department_closure(self) -> None:
        unit = "万千瓦时"
        rows = (
            Official2025Row(1, "总计", "总计", 100.0, unit, 1),
            Official2025Row(2, "部门合计", "部门合计", 100.0, unit, 1),
            Official2025Row(3, "父项", "父项", 70.0, unit, 1),
            Official2025Row(4, "新负荷", "新负荷", 10.0, unit, 1),
            Official2025Row(5, "排除", "排除", 5.0, unit, 1),
            Official2025Row(6, "其他", "其他", 30.0, unit, 1),
        )
        mapping = (
            OfficialMappingEntry(1, "总计", "汇总项", AccountingRole.SUMMARY),
            OfficialMappingEntry(2, "部门合计", "汇总项", AccountingRole.SUMMARY),
            OfficialMappingEntry(
                3, "父项", "传统及其他计量项", AccountingRole.SPLIT_PARENT, "部门"
            ),
            OfficialMappingEntry(
                4,
                "新负荷",
                "新型负荷",
                AccountingRole.SPLIT_COMPONENT,
                "部门",
                "父项",
            ),
            OfficialMappingEntry(
                5,
                "排除",
                "排除项",
                AccountingRole.SPLIT_COMPONENT,
                "部门",
                "父项",
            ),
            OfficialMappingEntry(
                6, "其他", "传统及其他计量项", AccountingRole.LEAF, "部门"
            ),
        )
        result = data_io.validate_official_balance(
            rows,
            mapping,
            total_label="总计",
            department_total_labels={"部门": "部门合计"},
        )
        self.assertTrue(result.passed)
        self.assertEqual(
            result.amount_by_classification,
            {"传统及其他计量项": 85.0, "排除项": 5.0, "新型负荷": 10.0},
        )

    def test_excluded_energy_ledger_is_itemized_unique_and_gross_net_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "official.xlsx"
            self._write_excluded_energy_fixture(source)
            ledger = data_io.build_excluded_energy_ledger(source)

        self.assertEqual(ledger.year.tolist(), list(range(2020, 2026)))
        self.assertTrue(ledger.year.is_unique)
        self.assertTrue(ledger.identity_passed.all())
        self.assertTrue(ledger.annual_special_deduction_applied.all())
        self.assertFalse(ledger.hourly_special_deduction_applied.any())
        self.assertFalse(ledger.special_items_in_correction_component.any())
        self.assertEqual(set(ledger.component_split_status), {"actual_itemized"})
        self.assertEqual(
            set(ledger.curve_energy_scope),
            {
                "natural_load_excluding_plant_aux_line_loss_"
                "pumped_storage_charging"
            },
        )
        row_2025 = ledger.loc[ledger.year.eq(2025)].iloc[0]
        self.assertEqual(row_2025.full_society_10k_kWh, 1_500.0)
        self.assertEqual(row_2025.plant_auxiliary_power_10k_kWh, 55.0)
        self.assertEqual(row_2025.line_loss_10k_kWh, 25.0)
        self.assertEqual(row_2025.pumped_storage_charging_10k_kWh, 15.0)
        self.assertEqual(row_2025.special_deduction_10k_kWh, 95.0)
        self.assertEqual(row_2025.natural_load_10k_kWh, 1_405.0)
        self.assertEqual(row_2025.natural_load_MWh, 14_050.0)
        self.assertAlmostEqual(row_2025.deduction_ratio, 95.0 / 1_500.0)
        self.assertEqual(len(row_2025.source_workbook_sha256), 64)

    def test_excluded_energy_ledger_rejects_missing_duplicate_negative_and_nonpositive_net(self) -> None:
        cases = (
            (
                {"omit_label": "线路损失电量"},
                "必须恰好出现一次",
            ),
            (
                {"duplicate_label": "线路损失电量"},
                "labels are not unique",
            ),
            (
                {"negative_label": "抽水蓄能抽水耗用电量"},
                "invalid 2020 value",
            ),
            (
                {"deduction_exceeds_total": True},
                "自然负荷电量必须为正",
            ),
        )
        for options, message in cases:
            with self.subTest(options=options), tempfile.TemporaryDirectory() as tmp:
                source = Path(tmp) / "official.xlsx"
                self._write_excluded_energy_fixture(source, **options)
                with self.assertRaisesRegex(ValueError, message):
                    data_io.build_excluded_energy_ledger(source)

    def _source_paths(self) -> tuple[Path, Path]:
        load = os.environ.get("V2_TEST_LOAD_WORKBOOK")
        official = os.environ.get("V2_TEST_OFFICIAL_WORKBOOK")
        if not load or not official:
            self.skipTest("real source workbook paths were not supplied")
        return Path(load), Path(official)

    def test_real_workbooks_mapping_balance_and_public_interfaces(self) -> None:
        load_path, official_path = self._source_paths()
        result = data_io.load_and_validate_inputs(
            load_workbook=load_path,
            official_workbook=official_path,
        )

        self.assertEqual(len(result["manifest"]), 4)
        self.assertTrue(all(len(item.sha256) == 64 for item in result["manifest"]))
        self.assertEqual(len(result["sample_catalog"]), 52)
        self.assertEqual(len({row.industry for row in result["sample_catalog"]}), 25)

        official = result["official_mapping"]
        required_columns = {
            "source_label",
            "department",
            "accounting_class",
            "electricity_2025",
            "is_leaf",
            "accounted_electricity_2025",
        }
        self.assertTrue(required_columns.issubset(official.columns))
        self.assertEqual(len(official), 135)
        self.assertTrue(official["source_row"].is_unique)

        balance = result["official_balance"]
        self.assertTrue(balance.passed)
        self.assertLessEqual(balance.absolute_error, balance.absolute_tolerance)
        self.assertLessEqual(balance.relative_error, 1e-10)
        self.assertTrue(all(item.passed for item in balance.department_balances))
        self.assertAlmostEqual(
            balance.amount_by_classification["排除项"], 7_812_227.3744, places=6
        )
        self.assertAlmostEqual(
            balance.amount_by_classification["新型负荷"], 866_414.3321, places=6
        )

        ledger = result["excluded_energy_ledger"]
        self.assertEqual(ledger.year.tolist(), list(range(2020, 2026)))
        actual_2025 = ledger.loc[ledger.year.eq(2025)].iloc[0]
        self.assertAlmostEqual(
            actual_2025.full_society_10k_kWh, 86_706_574.2129, places=5
        )
        self.assertAlmostEqual(
            actual_2025.special_deduction_10k_kWh, 7_812_227.3744, places=5
        )
        self.assertAlmostEqual(
            actual_2025.natural_load_10k_kWh, 78_894_346.8385, places=5
        )
        self.assertAlmostEqual(
            actual_2025.deduction_ratio, 0.09009959677587764, places=14
        )
        self.assertTrue(bool(actual_2025.identity_passed))

        label_energy = data_io.official_industry_energy(official)
        self.assertAlmostEqual(
            label_energy["全社会用电总计"], 86_706_574.2129, places=5
        )

        sample_energy = result["sample_industry_energy"]
        lookup = data_io.official_industry_energy(sample_energy)
        self.assertAlmostEqual(lookup["居民生活"], 11_256_824.8359, places=5)
        internet = sample_energy.loc[
            sample_energy["industry"].eq("互联网和相关服务")
        ].iloc[0]
        self.assertEqual(internet["official_energy_basis"], "split_parent_residual")
        self.assertAlmostEqual(
            internet["official_energy_2025"],
            302_614.7624 - 174_212.4337,
            places=6,
        )
        traditional = sample_energy.loc[sample_energy.is_traditional]
        self.assertEqual(len(traditional), 24)
        self.assertNotIn("run_group", traditional.columns)
        self.assertNotIn("parameter_group", traditional.columns)
        self.assertAlmostEqual(
            float(traditional.official_energy_2025.sum()),
            49_634_722.5846,
            places=5,
        )
        covered_share = float(traditional.official_energy_2025.sum()) / float(
            actual_2025.natural_load_10k_kWh
        )
        self.assertAlmostEqual(covered_share, 0.629129013339883, places=12)
        self.assertAlmostEqual(1.0 - covered_share, 0.370870986660117, places=12)

    def test_official_industry_energy_rejects_duplicate_keys(self) -> None:
        frame = pd.DataFrame(
            {
                "source_label": ["重复", "重复"],
                "electricity_2025": [1.0, 2.0],
            }
        )
        with self.assertRaisesRegex(ValueError, "unique"):
            data_io.official_industry_energy(frame)


if __name__ == "__main__":
    unittest.main()
