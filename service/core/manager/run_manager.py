#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
@Time    : 2026/10/21 15:40
@Author  : weiyutao
@File    : run_manager.py

运行业务组织:上传校验与保存、模板转 JSON、后台算法执行、结果提取。
算法包唯一调用方:只 import workstreams_load_characteristic_model_v2.pipeline.run_pipeline。
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import traceback
import uuid
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from core.exceptions import ConflictError, NotFoundError
from core.paths import (
    ANNUAL_ENERGY_MAX_BYTES,
    LOAD_WORKBOOK_MAX_BYTES,
    OFFICIAL_WORKBOOK_MAX_BYTES,
    RUNS_DIR,
)
from core.schema.run_dto import RunCreateDTO, RunStatusResp, RunUpdateDTO
from core.service.run_service import RunService
from core.table.run_record import RUN_STATUS_LABELS, RUN_STATUS_NAMES, RunRecord, RunStatus
from workstreams_load_characteristic_model_v2.pipeline import run_pipeline

HISTORY_SHEET_NAME = "历史负荷数据（2019-2023）"      # load_workbook 必含的全省历史曲线 sheet(括号为全角,与算法 config.py 一致)
TEMPLATE_SHEETS = ("annual_electricity", "scope_bridge")  # 年度电量模板必含的两张 sheet
PARAM_SCENARIOS = ("低碳", "基准", "高电气化")          # 规划情景允许取值(与算法包 SCENARIO_ORDER 一致)
DEFAULT_SCENARIO = "基准"                              # 情景缺省值(旧运行记录 params 无 scenario 键时兜底)
REPRESENTATIVE_TRAJECTORY = "P50_real"                 # 结果提取的目标代表轨迹
PARAM_YEARS = (2025, 2060)                             # 规划年允许区间(含端点)
PARAM_SAMPLES = (20, 1000)                             # 抽样数允许档位


@dataclass
class UploadedFile:
    """一次上传的单个文件。"""

    filename: str    # 原始文件名(含扩展名)
    content: bytes   # 文件内容


class RunManager:
    """运行组织:登记上传 → 后台执行算法 → 提取指定情景 P50 结果。"""

    def __init__(self, run_service: RunService) -> None:
        self.logger = logging.getLogger("RunManager")
        self.run_service = run_service
        # ponytail: 全局锁串行,一次只跑一个算法任务(CPU 密集,并发收益低);要并发时改任务队列按机器分发
        self._execute_lock = asyncio.Lock()

    async def create_run(
        self,
        load_workbook: UploadedFile,
        official_workbook: UploadedFile,
        annual_energy: UploadedFile | None,
        params: dict,
    ) -> str:
        """校验并保存上传文件,登记运行记录,返回 run_id。

        Args:
            load_workbook (UploadedFile): 用户日曲线工作簿(含全省历史曲线)。
            official_workbook (UploadedFile): 全社会用电分类表。
            annual_energy (UploadedFile | None): 年度电量模板,不传用包内默认。
            params (dict): 运行参数(start_year/end_year/sample_size/scenario)。

        Returns:
            str: 运行标识 run_id。
        """
        try:
            self._validate_params(params)
            self._check_file(load_workbook, LOAD_WORKBOOK_MAX_BYTES, "用户曲线工作簿")
            self._require_sheets(load_workbook, (HISTORY_SHEET_NAME,), "用户曲线工作簿")
            self._check_file(official_workbook, OFFICIAL_WORKBOOK_MAX_BYTES, "行业电量工作簿")
            if annual_energy is not None:
                self._check_file(annual_energy, ANNUAL_ENERGY_MAX_BYTES, "年度电量模板")
                self._require_sheets(annual_energy, TEMPLATE_SHEETS, "年度电量模板")

            run_id = uuid.uuid4().hex
            input_dir = RUNS_DIR / run_id / "inputs"
            output_dir = RUNS_DIR / run_id / "outputs"
            input_dir.mkdir(parents=True, exist_ok=True)
            output_dir.mkdir(parents=True, exist_ok=True)
            (input_dir / "load_workbook.xlsx").write_bytes(load_workbook.content)
            (input_dir / "official_workbook.xlsx").write_bytes(official_workbook.content)
            if annual_energy is not None:
                self._template_to_json(annual_energy, input_dir / "annual_energy_upstream.json")

            await self.run_service.create(
                RunCreateDTO(
                    run_id=run_id,
                    params=dict(
                        params,
                        filenames={
                            "load_workbook": load_workbook.filename,
                            "official_workbook": official_workbook.filename,
                            "annual_energy": annual_energy.filename if annual_energy else None,
                        },
                    ),
                    input_dir=str(input_dir),
                    output_dir=str(output_dir),
                )
            )
            self.logger.info("[RunManager]-%s- 运行已登记,输入目录 %s", run_id, input_dir)
            return run_id
        except ValueError as exc:
            self.logger.warning("[RunManager]-创建运行- 参数或文件校验未通过: %s", exc)
            raise
        except Exception as exc:
            self.logger.error("创建运行异常:\n%s", traceback.format_exc())
            raise RuntimeError(f"创建运行失败:{exc}") from exc

    async def execute_run(self, run_id: str) -> None:
        """后台执行算法并回写状态;由路由的 BackgroundTasks 调度,全局串行。

        Args:
            run_id (str): 运行标识。

        Returns: None
        """
        record = await self.run_service.get(run_id)
        if record is None:
            self.logger.error("[RunManager]-%s- 运行记录不存在,后台任务退出", run_id)
            return
        params = record.params
        scenario = params.get("scenario") or DEFAULT_SCENARIO
        input_dir = Path(record.input_dir)
        kwargs = dict(
            years=range(params["start_year"], params["end_year"] + 1),
            output_dir=record.output_dir,
            bootstrap_scenarios=params["sample_size"],
            planning_scenarios=[scenario],
            load_workbook=str(input_dir / "load_workbook.xlsx"),
            official_workbook=str(input_dir / "official_workbook.xlsx"),
            require_parquet=True,
        )
        upstream_json = input_dir / "annual_energy_upstream.json"
        if upstream_json.exists():
            kwargs["annual_energy_upstream"] = str(upstream_json)

        async with self._execute_lock:
            await self.run_service.update(run_id, RunUpdateDTO(status=RunStatus.RUNNING))
            try:
                await asyncio.to_thread(run_pipeline, **kwargs)
                result_path = await asyncio.to_thread(
                    self._extract_p50_csv, Path(record.output_dir), scenario
                )
                await self.run_service.update(
                    run_id,
                    RunUpdateDTO(status=RunStatus.SUCCEEDED, result_path=str(result_path)),
                )
                self.logger.info("[RunManager]-%s- 运行成功,结果 %s", run_id, result_path)
            except Exception as exc:
                self.logger.error("[RunManager]-%s- 算法执行失败:\n%s", run_id, traceback.format_exc())
                await self.run_service.update(
                    run_id, RunUpdateDTO(status=RunStatus.FAILED, error_message=str(exc))
                )

    async def get_status(self, run_id: str) -> RunStatusResp:
        """查询运行状态。

        Args:
            run_id (str): 运行标识。

        Returns:
            RunStatusResp: 状态响应结构。

        Raises:
            NotFoundError: 运行记录不存在。
        """
        try:
            record = await self._require_record(run_id)
            return RunStatusResp(
                job_id=record.run_id,
                status=RUN_STATUS_NAMES.get(record.status, "unknown"),
                error=record.error_message,
            )
        except ValueError as exc:
            self.logger.warning("[RunManager]-%s- 查询状态: %s", run_id, exc)
            raise
        except Exception as exc:
            self.logger.error("查询状态异常:\n%s", traceback.format_exc())
            raise RuntimeError(f"[RunManager]-{run_id}- 查询状态失败: {exc}") from exc

    async def get_result_path(self, run_id: str) -> Path:
        """返回结果 CSV 路径,供下载接口使用。

        Args:
            run_id (str): 运行标识。

        Returns:
            Path: 结果文件路径。

        Raises:
            NotFoundError: 运行记录不存在。
            ConflictError: 运行未成功,结果尚未生成。
        """
        record = await self._require_succeeded(run_id)
        return Path(record.result_path)

    async def get_result_views(self, run_id: str) -> dict:
        """查询运行结果的 JSON 视图,供前端绘图。

        结果按年拆组(views 每组一张图,格式对齐 regional_power_system_optimization 项目):
        每年曲线按本年最大标幺,跨年拼接无物理意义,前端分年渲染。

        Args:
            run_id (str): 运行标识。

        Returns:
            dict: {"views": [...]},每年一组 {year, sheet_name, index, data}。

        Raises:
            NotFoundError: 运行记录不存在。
            ConflictError: 运行未成功,结果尚未生成。
        """
        record = await self._require_succeeded(run_id)
        scenario = record.params.get("scenario") or DEFAULT_SCENARIO
        views = await asyncio.to_thread(
            self._build_result_views, Path(record.result_path), scenario
        )
        return {"views": views}

    async def _require_succeeded(self, run_id: str) -> RunRecord:
        """取已成功运行的记录,不存在/未成功即抛错。

        Args:
            run_id (str): 运行标识。

        Returns:
            RunRecord: 运行记录。

        Raises:
            NotFoundError: 运行记录不存在。
            ConflictError: 运行未成功,结果尚未生成。
        """
        record = await self._require_record(run_id)
        if record.status != RunStatus.SUCCEEDED or not record.result_path:
            raise ConflictError(
                f"运行 {run_id} 当前状态为 {RUN_STATUS_LABELS.get(record.status, '未知')},结果尚未生成"
            )
        return record

    async def _require_record(self, run_id: str) -> RunRecord:
        """取运行记录,不存在即抛 NotFoundError。

        Args:
            run_id (str): 运行标识。

        Returns:
            RunRecord: 运行记录。
        """
        record = await self.run_service.get(run_id)
        if record is None:
            raise NotFoundError(f"未找到运行记录:{run_id}")
        return record

    @staticmethod
    def _validate_params(params: dict) -> None:
        """参数校验:年份区间、抽样档位与情景取值,错报指明字段名。

        Args:
            params (dict): 运行参数(start_year/end_year/sample_size/scenario)。

        Returns: None
        """
        start = params.get("start_year")
        end = params.get("end_year")
        sample = params.get("sample_size")
        if not isinstance(start, int) or not isinstance(end, int):
            raise ValueError("参数 start_year/end_year 必须为整数")
        if not (PARAM_YEARS[0] <= start <= end <= PARAM_YEARS[1]):
            raise ValueError(
                f"参数年份区间 {start}—{end} 超出允许范围 {PARAM_YEARS[0]}—{PARAM_YEARS[1]}"
            )
        if sample not in PARAM_SAMPLES:
            raise ValueError(f"参数 sample_size={sample} 不在允许档位 {PARAM_SAMPLES}")
        if params.get("scenario", DEFAULT_SCENARIO) not in PARAM_SCENARIOS:
            raise ValueError(
                f"参数 scenario={params.get('scenario')} 不在允许取值 {PARAM_SCENARIOS}"
            )

    @staticmethod
    def _check_file(uploaded: UploadedFile, max_bytes: int, kind: str) -> None:
        """传输层校验:扩展名、非空与大小上限。

        Args:
            uploaded (UploadedFile): 上传文件。
            max_bytes (int): 大小上限(字节)。
            kind (str): 文件类别中文名,用于错误信息。

        Returns: None
        """
        if not uploaded.filename.lower().endswith(".xlsx"):
            raise ValueError(f"{kind}文件 {uploaded.filename} 扩展名不是 .xlsx")
        if not uploaded.content:
            raise ValueError(f"{kind}文件 {uploaded.filename} 内容为空")
        if len(uploaded.content) > max_bytes:
            raise ValueError(
                f"{kind}文件 {uploaded.filename} 大小 {len(uploaded.content)} 字节,"
                f"超过上限 {max_bytes} 字节"
            )

    @staticmethod
    def _require_sheets(uploaded: UploadedFile, required: tuple[str, ...], kind: str) -> None:
        """结构快检:工作簿必须包含指定 sheet。

        Args:
            uploaded (UploadedFile): 上传文件。
            required (tuple[str, ...]): 必含 sheet 名。
            kind (str): 文件类别中文名,用于错误信息。

        Returns: None
        """
        workbook = openpyxl.load_workbook(io.BytesIO(uploaded.content), read_only=True)
        names = set(workbook.sheetnames)
        workbook.close()
        missing = [name for name in required if name not in names]
        if missing:
            raise ValueError(f"{kind}文件 {uploaded.filename} 缺少 sheet:{missing}")

    @staticmethod
    def _template_to_json(uploaded: UploadedFile, out_path: Path) -> None:
        """年度电量模板两 sheet 转上游 JSON(键与 sheet 名一致,列名即记录字段名)。

        Args:
            uploaded (UploadedFile): 模板文件。
            out_path (Path): 输出 JSON 路径。

        Returns: None
        """
        frames = pd.read_excel(io.BytesIO(uploaded.content), sheet_name=list(TEMPLATE_SHEETS))
        payload: dict = {}
        for sheet in TEMPLATE_SHEETS:
            frame = frames[sheet]
            if frame.empty:
                raise ValueError(f"年度电量模板 sheet {sheet} 无数据行")
            na_columns = frame.columns[frame.isna().any()].tolist()
            if na_columns:
                raise ValueError(f"年度电量模板 sheet {sheet} 存在空值,涉及列:{na_columns}")
            payload[sheet] = frame.to_dict(orient="records")
        out_path.write_text(
            json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding="utf-8"
        )

    @staticmethod
    def _extract_p50_csv(output_dir: Path, scenario: str) -> Path:
        """从算法输出提取指定情景 P50 代表轨迹的两列 CSV(timestamp, load_pu)。

        Args:
            output_dir (Path): 算法输出目录。
            scenario (str): 结果提取的目标情景。

        Returns:
            Path: 结果 CSV 路径。

        Raises:
            ValueError: 输出缺少逐时曲线文件或目标轨迹切片为空。
        """
        parquet = output_dir / "逐时曲线" / "province_hourly_normalized_trajectories.parquet"
        if not parquet.exists():
            raise ValueError(f"算法输出缺少逐时曲线文件:{parquet}")
        frame = pd.read_parquet(
            parquet, columns=["timestamp", "planning_scenario", "trajectory", "load_pu"]
        )
        picked = frame[
            (frame["planning_scenario"] == scenario)
            & (frame["trajectory"] == REPRESENTATIVE_TRAJECTORY)
        ]
        if picked.empty:
            raise ValueError(f"逐时曲线中未找到{scenario}情景 P50_real 轨迹")
        result_path = output_dir.parent / "result" / f"{scenario}_p50_hourly.csv"
        result_path.parent.mkdir(parents=True, exist_ok=True)
        picked[["timestamp", "load_pu"]].to_csv(result_path, index=False)
        return result_path

    @staticmethod
    def _build_result_views(csv_path: Path, scenario: str) -> list[dict]:
        """结果 CSV 转按年分组的绘图视图(float_precision 取 round_trip,值与 CSV 逐位一致)。

        Args:
            csv_path (Path): 结果 CSV 路径。
            scenario (str): 结果对应的规划情景,用于 sheet_name 标签。

        Returns:
            list[dict]: 每年一组 {year, sheet_name, index, data}(data 为列数组,单列 load_pu)。
        """
        frame = pd.read_csv(csv_path, dtype={"timestamp": str}, float_precision="round_trip")
        views: list[dict] = []
        for year, rows in frame.groupby(frame["timestamp"].str[:4]):
            views.append(
                {
                    "year": int(year),
                    "sheet_name": f"{scenario}_p50_{year}",
                    "index": rows["timestamp"].tolist(),
                    "data": [{"column_name": "load_pu", "data": rows["load_pu"].tolist()}],
                }
            )
        return views
