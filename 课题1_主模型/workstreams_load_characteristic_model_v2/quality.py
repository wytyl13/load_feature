"""读取并清洗全省与大用户日负荷。"""

from __future__ import annotations

import hashlib
import itertools
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .calendar_utils import classify_date, season_of
from .config import INDUSTRY_AGGREGATION_MODE, SHAPE_NUMERICAL_TOLERANCE


HOUR_COLS = [f"h{hour:02d}" for hour in range(24)]
SHAPE_COLS = [f"p{hour:02d}" for hour in range(24)]

_INDUSTRY_DAY_KEYS = [
    "department",
    "industry",
    "date",
    "year",
    "season",
    "day_type_detail",
]
_INDUSTRY_DAY_QUALITY_COLS = [
    *_INDUSTRY_DAY_KEYS,
    "source_user_count",
    "active_user_count",
    "missing_user_count",
    "missing_user_rate",
    "partial_composition_day",
    "active_user_signature",
    "shape_sum_error",
]


def _empty_industry_daily() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            *_INDUSTRY_DAY_KEYS,
            "day_type",
            "day_group",
            "source_user_count",
            "active_user_count",
            "missing_user_count",
            "missing_user_rate",
            "partial_composition_day",
            "active_user_signature",
            "aggregation_method",
            *SHAPE_COLS,
        ]
    )


def _empty_industry_daily_diagnostics() -> pd.DataFrame:
    return pd.DataFrame(columns=_INDUSTRY_DAY_QUALITY_COLS)


def aggregate_industry_daily(
    user_daily: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Aggregate normalized valid user-days into equal-weight industry-days.

    Every row is already expected to be a valid, unit-sum user-day (the
    contract of :func:`read_user_daily`).  A user absent on a date contributes
    no row and is never represented by a zero curve.  Composition changes are
    retained as diagnostics only.
    """

    if user_daily.empty:
        return _empty_industry_daily(), _empty_industry_daily_diagnostics()

    frame = user_daily.copy()
    if "department" not in frame.columns and "broad_sector" in frame.columns:
        frame["department"] = frame["broad_sector"]
    if "user" not in frame.columns and "user_key" in frame.columns:
        frame["user"] = frame["user_key"]
    if "day_type_detail" not in frame.columns:
        if "day_type" in frame.columns:
            frame["day_type_detail"] = frame["day_type"]
        elif "day_group" in frame.columns:
            frame["day_type_detail"] = frame["day_group"]

    required = {*_INDUSTRY_DAY_KEYS, "user", *SHAPE_COLS}
    missing = sorted(required.difference(frame.columns))
    if missing:
        raise ValueError(f"用户日数据缺少字段：{missing}")
    metadata_columns = [*_INDUSTRY_DAY_KEYS, "user"]
    if frame[metadata_columns].isna().any().any():
        raise ValueError("用户日数据的行业、用户、日期或日历标签存在缺失")

    frame["date"] = pd.to_datetime(frame["date"], errors="raise").dt.normalize()
    if not frame.year.astype(int).eq(frame.date.dt.year).all():
        raise ValueError("用户日year与date年份不一致")
    if frame[list(SHAPE_COLS)].isna().any().any():
        raise ValueError("行业日聚合输入存在缺失形状值")
    duplicate_keys = ["department", "industry", "date", "user"]
    if frame.duplicated(duplicate_keys).any():
        duplicate = frame.loc[
            frame.duplicated(duplicate_keys, keep=False), duplicate_keys
        ].head(10)
        raise ValueError(f"用户日记录重复：{duplicate.to_dict('records')}")

    # A natural day may not acquire conflicting calendar metadata merely by
    # splitting users into separate groups.
    metadata = ["year", "season", "day_type_detail"]
    conflicts = (
        frame.groupby(["department", "industry", "date"], observed=True)[metadata]
        .nunique(dropna=False)
        .gt(1)
        .any(axis=1)
    )
    if conflicts.any():
        raise ValueError(
            "同一行业自然日存在冲突的年份、季节或日型："
            f"{conflicts.loc[conflicts].index.tolist()[:10]}"
        )

    source_users = (
        frame.groupby(["department", "industry"], observed=True)["user"]
        .nunique()
        .rename("source_user_count")
    )
    rows: list[dict[str, object]] = []
    diagnostics: list[dict[str, object]] = []
    for key, part in frame.groupby(_INDUSTRY_DAY_KEYS, observed=True, sort=True):
        values = part[SHAPE_COLS].to_numpy(dtype=float)
        if not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"行业日聚合输入存在非法形状值：{key}")
        sums = values.sum(axis=1)
        if not np.allclose(
            sums,
            1.0,
            rtol=0.0,
            atol=SHAPE_NUMERICAL_TOLERANCE,
        ):
            raise ValueError(f"用户日未完成归一化：{key}")

        curve = values.mean(axis=0)
        curve_sum = float(curve.sum())
        if abs(curve_sum - 1.0) > SHAPE_NUMERICAL_TOLERANCE:
            raise RuntimeError(f"行业日等权聚合未守恒：{key}")
        curve = curve / curve_sum

        department, industry, date, year, season, day_type_detail = key
        total_users = int(source_users.loc[(department, industry)])
        active_users = tuple(sorted(part["user"].astype(str).unique()))
        active_count = len(active_users)
        missing_count = total_users - active_count
        signature = hashlib.sha1("\x1f".join(active_users).encode("utf-8")).hexdigest()
        common = {
            "department": department,
            "industry": industry,
            "date": pd.Timestamp(date),
            "year": int(year),
            "season": season,
            "day_type_detail": day_type_detail,
            "source_user_count": total_users,
            "active_user_count": active_count,
            "missing_user_count": missing_count,
            "missing_user_rate": missing_count / total_users,
            "partial_composition_day": active_count < total_users,
            "active_user_signature": signature,
        }
        rows.append(
            {
                **common,
                # Keep the existing calendar aliases available to generic
                # readers while making the detailed label authoritative.
                "day_type": day_type_detail,
                "day_group": (
                    "工作日"
                    if "工作日" in str(day_type_detail)
                    and "非工作日" not in str(day_type_detail)
                    else "非工作日"
                ),
                "aggregation_method": INDUSTRY_AGGREGATION_MODE,
                **dict(zip(SHAPE_COLS, curve, strict=True)),
            }
        )
        diagnostics.append(
            {
                **common,
                "shape_sum_error": abs(float(curve.sum()) - 1.0),
            }
        )

    daily = pd.DataFrame(rows).sort_values(
        ["department", "industry", "date"], kind="mergesort"
    ).reset_index(drop=True)
    diagnostic_frame = pd.DataFrame(diagnostics).sort_values(
        ["department", "industry", "date"], kind="mergesort"
    ).reset_index(drop=True)
    return daily, diagnostic_frame


def parse_user_sheet_name(name: str) -> tuple[str, str, str]:
    parts = str(name).strip().split("-", 2)
    if len(parts) != 3 or any(not part.strip() for part in parts):
        raise ValueError(f"样本工作表名称不符合 部门-细行业-用户：{name!r}")
    return tuple(part.strip() for part in parts)  # type: ignore[return-value]


def _curve_digest(values: np.ndarray) -> str:
    return hashlib.sha1(np.asarray(values, dtype=np.float64).tobytes()).hexdigest()


def _systemic_copy_marks(frame: pd.DataFrame, min_users: int = 10) -> set[tuple[str, pd.Timestamp]]:
    """识别同一短间隔日期对在至少min_users个用户中同步重复的原始曲线。"""
    pair_users: dict[tuple[pd.Timestamp, pd.Timestamp], set[str]] = defaultdict(set)
    for user_key, user_frame in frame.groupby("user_key", sort=False):
        for _, repeated in user_frame.groupby("curve_digest", sort=False):
            dates = sorted(set(pd.to_datetime(repeated.date)))
            if len(dates) < 2:
                continue
            for first, second in itertools.combinations(dates, 2):
                if 0 < (second - first).days <= 31:
                    pair_users[(first, second)].add(str(user_key))
    marked: set[tuple[str, pd.Timestamp]] = set()
    for (first, second), users in pair_users.items():
        if len(users) < min_users:
            continue
        for user_key in users:
            marked.add((user_key, first))
            marked.add((user_key, second))
    return marked


def read_system_history(path: str | Path) -> pd.DataFrame:
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    required = "历史负荷数据（2019-2023）"
    if required not in workbook.sheetnames:
        raise ValueError(f"缺少工作表：{required}")
    sheet = workbook[required]
    rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        if values[0] is None:
            continue
        hourly = pd.to_numeric(pd.Series(values[1:25]), errors="coerce").to_numpy(float)
        stamp = pd.Timestamp(values[0]).normalize()
        rows.append({"date": stamp, **dict(zip(HOUR_COLS, hourly))})
    frame = pd.DataFrame(rows).sort_values("date").reset_index(drop=True)
    if frame.empty or frame.date.duplicated().any():
        raise ValueError("全省历史负荷为空或日期重复")
    return frame


def read_user_daily(path: str | Path, include_new_load: bool = False) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    records: list[dict] = []
    quality: list[dict] = []
    for sheet_name in workbook.sheetnames[2:]:
        department, industry, user = parse_user_sheet_name(sheet_name)
        source_industry = industry
        if department == "居民生活":
            industry = "居民生活"
            user = f"{source_industry}-{user}"
        is_new = "充换电" in industry
        if is_new and not include_new_load:
            quality.append({"sheet": sheet_name.strip(), "department": department, "industry": industry, "user": user, "status": "留待后续新型负荷模块（不进入本轮24行业形状；电量未扣除）", "total_days": 0, "valid_days_pre_copy_check": 0})
            continue
        sheet = workbook[sheet_name]
        counts: dict[str, int] = {}
        total_days = 0
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if values[0] is None:
                continue
            total_days += 1
            stamp = pd.Timestamp(values[0]).normalize()
            hourly = pd.to_numeric(pd.Series(values[1:25]), errors="coerce").to_numpy(float)
            reason = "有效"
            if len(hourly) != 24 or not np.isfinite(hourly).all():
                reason = "缺失或非有限值"
            elif np.any(hourly < 0):
                reason = "负值"
            elif float(hourly.sum()) <= 0:
                reason = "日电量非正"
            elif float(np.max(hourly) - np.min(hourly)) <= 1e-12:
                reason = "24小时平线"
            counts[reason] = counts.get(reason, 0) + 1
            if reason != "有效":
                continue
            energy = float(hourly.sum())
            detailed, day_group, calendar_source = classify_date(stamp)
            normalized = hourly / energy
            records.append({
                "date": stamp,
                "year": stamp.year,
                "month": stamp.month,
                "season": season_of(stamp.month),
                "day_type": detailed,
                "day_group": day_group,
                "calendar_source": calendar_source,
                "department": department,
                "industry": industry,
                "source_industry": source_industry,
                "user": user,
                "user_key": f"{industry}-{user}",
                "sample_energy": energy,
                "curve_digest": _curve_digest(hourly),
                **dict(zip(HOUR_COLS, hourly)),
                **dict(zip(SHAPE_COLS, normalized)),
            })
        quality.append({
            "sheet": sheet_name.strip(),
            "department": department,
            "industry": industry,
            "source_industry": source_industry,
            "user": user,
            "status": "已读取",
            "total_days": total_days,
            "valid_days_pre_copy_check": counts.get("有效", 0),
            "missing_or_nonfinite": counts.get("缺失或非有限值", 0),
            "negative": counts.get("负值", 0),
            "nonpositive": counts.get("日电量非正", 0),
            "flat": counts.get("24小时平线", 0),
        })
    frame = pd.DataFrame(records)
    if frame.empty:
        raise ValueError("没有可用的大用户形状日")
    marked = _systemic_copy_marks(frame)
    synchronized = pd.Series(
        [(str(user_key), pd.Timestamp(stamp)) in marked for user_key, stamp in zip(frame.user_key, frame.date)],
        index=frame.index,
    )
    excluded = frame.loc[synchronized].groupby(["department", "industry", "user"]).size().rename("synchronized_copy_days")
    frame = frame.loc[~synchronized].drop(columns="curve_digest").reset_index(drop=True)
    qframe = pd.DataFrame(quality)
    if not excluded.empty:
        qframe = qframe.merge(excluded.reset_index(), how="left", on=["department", "industry", "user"])
    if "synchronized_copy_days" not in qframe:
        qframe["synchronized_copy_days"] = 0
    else:
        qframe["synchronized_copy_days"] = qframe["synchronized_copy_days"].fillna(0).astype(int)
    qframe = qframe.drop(columns=["valid_days"], errors="ignore")
    valid_counts = frame.groupby(["department", "industry", "user"]).size().rename("valid_days")
    qframe = qframe.merge(valid_counts.reset_index(), how="left", on=["department", "industry", "user"])
    qframe["valid_days"] = qframe["valid_days"].fillna(0).astype(int)
    qframe["valid_rate"] = np.where(qframe.total_days > 0, qframe.valid_days / qframe.total_days, 0.0)
    return frame, qframe
