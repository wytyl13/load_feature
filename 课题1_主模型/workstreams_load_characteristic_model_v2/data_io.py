"""Input inventory, schema checks, and official-industry electricity access.

The formal runtime reads only versioned snapshots under
``负荷特性模型（完善版）/inputs``.  Every reader also accepts an explicit path so
the preparation step and tests can validate a source before it is snapshotted.
No reader silently searches a personal Desktop or chat-download directory.
"""

from __future__ import annotations

import csv
import hashlib
import math
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import pandas as pd
from openpyxl import load_workbook

try:  # package import
    from .config import (
        DEFAULT_LOAD_WORKBOOK,
        DEFAULT_OFFICIAL_WORKBOOK,
        OFFICIAL_2025_MAPPING_CSV,
        OFFICIAL_DEPARTMENT_TOTAL_LABELS,
        OFFICIAL_EXCLUDED_ENERGY_ACTUAL_YEARS,
        OFFICIAL_EXCLUDED_ENERGY_LABELS,
        OFFICIAL_EXPECTED_UNIT,
        HISTORICAL_SHAPE_SCOPE,
        NATURAL_LOAD_CURVE_ENERGY_SCOPE,
        OFFICIAL_RELATIVE_TOLERANCE,
        OFFICIAL_TOTAL_LABEL,
        RESERVED_LOAD_SHEETS,
        SAMPLE_INDUSTRY_MAPPING_CSV,
        TRADITIONAL_COMPONENT_COUNT,
    )
    from .schemas import (
        AccountingRole,
        DepartmentBalance,
        InputFileManifest,
        Official2025Row,
        OfficialBalanceResult,
        OfficialMappingEntry,
        SampleIndustryEnergy,
        SampleIndustryMapping,
        SampleSheetRecord,
    )
except ImportError:  # direct script/test import
    from config import (  # type: ignore
        DEFAULT_LOAD_WORKBOOK,
        DEFAULT_OFFICIAL_WORKBOOK,
        OFFICIAL_2025_MAPPING_CSV,
        OFFICIAL_DEPARTMENT_TOTAL_LABELS,
        OFFICIAL_EXCLUDED_ENERGY_ACTUAL_YEARS,
        OFFICIAL_EXCLUDED_ENERGY_LABELS,
        OFFICIAL_EXPECTED_UNIT,
        HISTORICAL_SHAPE_SCOPE,
        NATURAL_LOAD_CURVE_ENERGY_SCOPE,
        OFFICIAL_RELATIVE_TOLERANCE,
        OFFICIAL_TOTAL_LABEL,
        RESERVED_LOAD_SHEETS,
        SAMPLE_INDUSTRY_MAPPING_CSV,
        TRADITIONAL_COMPONENT_COUNT,
    )
    from schemas import (  # type: ignore
        AccountingRole,
        DepartmentBalance,
        InputFileManifest,
        Official2025Row,
        OfficialBalanceResult,
        OfficialMappingEntry,
        SampleIndustryEnergy,
        SampleIndustryMapping,
        SampleSheetRecord,
    )


_EDGE_WHITESPACE = " \t\r\n\u00a0\u3000"
_USER_SHEET_RE = re.compile(r"^(?P<broad>.+?)-(?P<industry>.+?)-用户(?P<number>\d+)$")
_BOOL_TRUE = {"1", "true", "yes", "y", "是"}
_BOOL_FALSE = {"0", "false", "no", "n", "否", ""}


def clean_text(value: object) -> str:
    """Return text with only leading/trailing whitespace removed.

    Internal whitespace and Chinese punctuation are preserved because they are
    meaningful parts of official source labels.
    """

    if value is None:
        return ""
    return str(value).strip(_EDGE_WHITESPACE)


def _parse_bool(value: object, *, field: str) -> bool:
    token = clean_text(value).lower()
    if token in _BOOL_TRUE:
        return True
    if token in _BOOL_FALSE:
        return False
    raise ValueError(f"{field} must be a boolean token, got {value!r}")


def sha256_file(path: str | Path, *, chunk_size: int = 1024 * 1024) -> str:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_input_manifest(paths: Mapping[str, str | Path]) -> tuple[InputFileManifest, ...]:
    """Build a deterministic SHA256 manifest without copying any input."""

    records: list[InputFileManifest] = []
    for logical_name, raw_path in sorted(paths.items()):
        path = Path(raw_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"missing required input {logical_name}: {path}")
        stat = path.stat()
        modified = datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()
        records.append(
            InputFileManifest(
                logical_name=clean_text(logical_name),
                path=path,
                size_bytes=int(stat.st_size),
                modified_time_utc=modified,
                sha256=sha256_file(path),
            )
        )
    return tuple(records)


def verify_input_manifest(
    manifest: Iterable[InputFileManifest],
) -> None:
    """Raise if a file no longer matches its recorded size or SHA256."""

    for item in manifest:
        if not item.path.is_file():
            raise FileNotFoundError(f"manifest input disappeared: {item.path}")
        if item.path.stat().st_size != item.size_bytes:
            raise ValueError(f"input size changed after manifesting: {item.logical_name}")
        current = sha256_file(item.path)
        if current != item.sha256:
            raise ValueError(f"input SHA256 changed after manifesting: {item.logical_name}")


def parse_sample_sheet_name(sheet_name: object) -> SampleSheetRecord:
    raw = "" if sheet_name is None else str(sheet_name)
    cleaned = clean_text(raw)
    match = _USER_SHEET_RE.fullmatch(cleaned)
    if not match:
        raise ValueError(f"invalid sample sheet name: {raw!r}")
    broad_sector = clean_text(match.group("broad"))
    industry = clean_text(match.group("industry"))
    if not broad_sector or not industry:
        raise ValueError(f"empty sector or industry in sample sheet: {raw!r}")
    # Residential sheets use the middle token for a locality (currently 市中),
    # while the approved fine-industry mapping treats residential consumption
    # as one province-wide industry.  Preserve the raw sheet name for audit and
    # normalize only the industry key consumed by the model.
    if broad_sector == "居民生活":
        industry = "居民生活"
    return SampleSheetRecord(
        raw_sheet_name=raw,
        sheet_name=cleaned,
        broad_sector=broad_sector,
        industry=industry,
        user_number=int(match.group("number")),
    )


def read_sample_sheet_catalog(
    workbook_path: str | Path = DEFAULT_LOAD_WORKBOOK,
) -> tuple[SampleSheetRecord, ...]:
    """Read and parse all user sheets, cleaning the known leading-space issue."""

    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        cleaned_titles = [clean_text(name) for name in workbook.sheetnames]
        missing_reserved = sorted(set(RESERVED_LOAD_SHEETS).difference(cleaned_titles))
        if missing_reserved:
            raise ValueError(f"load workbook missing reserved sheets: {missing_reserved}")
        records = [
            parse_sample_sheet_name(raw_name)
            for raw_name, clean_name in zip(workbook.sheetnames, cleaned_titles)
            if clean_name not in RESERVED_LOAD_SHEETS
        ]
    finally:
        workbook.close()
    keys = [(row.industry, row.user_number) for row in records]
    if len(keys) != len(set(keys)):
        duplicates = sorted({key for key in keys if keys.count(key) > 1})
        raise ValueError(f"duplicate industry/user sample sheets: {duplicates}")
    return tuple(records)


def _display_decimal_places(number_format: object) -> int:
    fmt = clean_text(number_format).split(";", 1)[0]
    match = re.search(r"[0#]+\.([0#]+)", fmt)
    return len(match.group(1)) if match else 0


def read_official_year_rows(
    workbook_path: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
    *,
    year: int,
) -> tuple[Official2025Row, ...]:
    """Read every labelled row for one official workbook year.

    ``Official2025Row.value_2025`` is retained as a compatibility field name;
    when this generic reader is used it contains the requested year's value.
    """

    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        year_header_row = None
        year_column = None
        # Prefer the explicit table-header token (for example ``2020年``).
        # Some source files contain a bare year in the metadata block above the
        # table, which must not be mistaken for the electricity-data column.
        for expected_token in (f"{int(year)}年", str(int(year))):
            for row in range(1, min(worksheet.max_row, 40) + 1):
                for column in range(1, min(worksheet.max_column, 20) + 1):
                    token = clean_text(worksheet.cell(row, column).value)
                    if token == expected_token:
                        year_header_row, year_column = row, column
                        break
                if year_header_row is not None:
                    break
            if year_header_row is not None:
                break
        if year_header_row is None or year_column is None:
            raise ValueError(f"official workbook has no {year} year header")

        unit = ""
        for row in range(1, year_header_row + 1):
            for column in range(1, min(worksheet.max_column, 8) + 1):
                token = clean_text(worksheet.cell(row, column).value)
                if "单位" in token:
                    unit = clean_text(re.split(r"[:：]", token, maxsplit=1)[-1])
                    break
            if unit:
                break
        if not unit:
            raise ValueError("official workbook unit is missing")
        if unit != OFFICIAL_EXPECTED_UNIT:
            raise ValueError(
                f"official workbook unit must be {OFFICIAL_EXPECTED_UNIT!r}, got {unit!r}"
            )

        records: list[Official2025Row] = []
        for row in range(year_header_row + 1, worksheet.max_row + 1):
            raw_label = worksheet.cell(row, 1).value
            label = clean_text(raw_label)
            if not label:
                continue
            cell = worksheet.cell(row, year_column)
            value = cell.value
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"official row {row} has non-numeric {year} value: {value!r}")
            numeric = float(value)
            if not math.isfinite(numeric) or numeric < 0:
                raise ValueError(f"official row {row} has invalid {year} value: {numeric}")
            records.append(
                Official2025Row(
                    source_row=row,
                    raw_label="" if raw_label is None else str(raw_label),
                    label=label,
                    value_2025=numeric,
                    unit=unit,
                    display_decimal_places=_display_decimal_places(cell.number_format),
                )
            )
    finally:
        workbook.close()

    labels = [row.label for row in records]
    if len(labels) != len(set(labels)):
        duplicates = sorted({label for label in labels if labels.count(label) > 1})
        raise ValueError(f"official {int(year)} labels are not unique: {duplicates}")
    return tuple(records)


def read_official_2025_rows(
    workbook_path: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
) -> tuple[Official2025Row, ...]:
    """Compatibility wrapper for the formal 2025 mapping gate."""

    return read_official_year_rows(workbook_path, year=2025)


def build_excluded_energy_ledger(
    workbook_path: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
) -> pd.DataFrame:
    """Build the audited 2020—2025 plant-use/loss/pumping exclusion ledger.

    The official source unit is 10,000 kWh, equal to 10 MWh.  Each required
    source label must occur exactly once in every year.  Natural-load energy is
    derived only as ``full society - the three itemized exclusions``; the
    function validates the item sum, the gross-to-net bridge, units, signs and
    year coverage before returning one row per official actual year.
    """

    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    years = tuple(int(year) for year in OFFICIAL_EXCLUDED_ENERGY_ACTUAL_YEARS)
    if not years or len(years) != len(set(years)):
        raise AssertionError("官方三项排除电量年份配置必须非空且唯一")
    if years != tuple(range(2020, 2026)):
        raise AssertionError("官方三项排除电量实值期必须严格为2020—2025年")
    if len(OFFICIAL_EXCLUDED_ENERGY_LABELS) != 3 or len(
        set(OFFICIAL_EXCLUDED_ENERGY_LABELS.values())
    ) != 3:
        raise AssertionError("厂用电、线损和抽蓄用电源标签必须恰有三个且互不重复")

    resolved_path = path.resolve()
    source_sha256 = sha256_file(resolved_path)
    rows: list[dict[str, object]] = []
    required_labels = {
        "full_society": OFFICIAL_TOTAL_LABEL,
        **OFFICIAL_EXCLUDED_ENERGY_LABELS,
    }
    for year in years:
        official_rows = read_official_year_rows(resolved_path, year=year)
        matches: dict[str, Official2025Row] = {}
        for field, label in required_labels.items():
            found = [row for row in official_rows if row.label == label]
            if len(found) != 1:
                raise ValueError(
                    f"官方{year}年源标签必须恰好出现一次：{label!r}，实际{len(found)}次"
                )
            matches[field] = found[0]

        units = {row.unit for row in matches.values()}
        if units != {OFFICIAL_EXPECTED_UNIT}:
            raise ValueError(
                f"官方{year}年三项排除桥接单位不一致：{sorted(units)}"
            )
        values = {
            field: float(row.value_2025) for field, row in matches.items()
        }
        invalid = {
            field: value
            for field, value in values.items()
            if not math.isfinite(value) or value < 0
        }
        if invalid:
            raise ValueError(f"官方{year}年三项排除桥接含非有限值或负值：{invalid}")

        full = values["full_society"]
        plant_aux = values["plant_auxiliary_power"]
        line_loss = values["line_loss"]
        pumped_storage = values["pumped_storage_charging"]
        deduction = plant_aux + line_loss + pumped_storage
        natural = full - deduction
        if full <= 0:
            raise ValueError(f"官方{year}年全社会电量必须为正")
        if deduction <= 0:
            raise ValueError(f"官方{year}年三项排除电量合计必须为正")
        if natural <= 0:
            raise ValueError(
                f"官方{year}年自然负荷电量必须为正：Q_full={full}, D={deduction}"
            )

        max_decimals = max(row.display_decimal_places for row in matches.values())
        absolute_tolerance = 0.5 * 10.0 ** (-max_decimals)
        item_sum_error = abs(deduction - (plant_aux + line_loss + pumped_storage))
        bridge_error = abs(full - (natural + deduction))
        bridge_relative_error = bridge_error / full
        identity_passed = bool(
            item_sum_error <= absolute_tolerance
            and bridge_error <= absolute_tolerance
            and bridge_relative_error <= OFFICIAL_RELATIVE_TOLERANCE
        )
        if not identity_passed:
            raise ValueError(
                f"官方{year}年三项排除桥接恒等式失败："
                f"item_sum_error={item_sum_error:.12g}, "
                f"bridge_error={bridge_error:.12g}"
            )

        record: dict[str, object] = {
            "year": year,
            "full_society_10k_kWh": full,
            "plant_auxiliary_power_10k_kWh": plant_aux,
            "line_loss_10k_kWh": line_loss,
            "pumped_storage_charging_10k_kWh": pumped_storage,
            "special_deduction_10k_kWh": deduction,
            "natural_load_10k_kWh": natural,
            "deduction_ratio": deduction / full,
            "item_sum_absolute_error_10k_kWh": item_sum_error,
            "bridge_absolute_error_10k_kWh": bridge_error,
            "bridge_relative_error": bridge_relative_error,
            "identity_tolerance_10k_kWh": absolute_tolerance,
            "identity_passed": identity_passed,
            "source_type": "official_actual_itemized",
            "source_workbook": str(resolved_path),
            "source_workbook_sha256": source_sha256,
            "source_unit": OFFICIAL_EXPECTED_UNIT,
            "full_society_source_label": matches["full_society"].label,
            "plant_auxiliary_power_source_label": matches[
                "plant_auxiliary_power"
            ].label,
            "line_loss_source_label": matches["line_loss"].label,
            "pumped_storage_charging_source_label": matches[
                "pumped_storage_charging"
            ].label,
            "full_society_source_row": matches["full_society"].source_row,
            "plant_auxiliary_power_source_row": matches[
                "plant_auxiliary_power"
            ].source_row,
            "line_loss_source_row": matches["line_loss"].source_row,
            "pumped_storage_charging_source_row": matches[
                "pumped_storage_charging"
            ].source_row,
            "component_split_status": "actual_itemized",
            "annual_special_deduction_applied": True,
            "future_special_item_split_available": False,
            "hourly_special_deduction_applied": False,
            "special_items_in_correction_component": False,
            "curve_energy_scope": NATURAL_LOAD_CURVE_ENERGY_SCOPE,
            "historical_shape_scope": HISTORICAL_SHAPE_SCOPE,
        }
        for stem in (
            "full_society",
            "plant_auxiliary_power",
            "line_loss",
            "pumped_storage_charging",
            "special_deduction",
            "natural_load",
        ):
            record[f"{stem}_MWh"] = float(record[f"{stem}_10k_kWh"]) * 10.0
        rows.append(record)

    ledger = pd.DataFrame.from_records(rows).sort_values("year").reset_index(drop=True)
    if len(ledger) != len(years) or ledger.year.duplicated().any():
        raise AssertionError("官方三项排除电量台账必须严格为2020—2025每年一行")
    if tuple(ledger.year.astype(int)) != years:
        raise AssertionError("官方三项排除电量台账年份覆盖或顺序不完整")
    numeric_columns = [
        column
        for column in ledger.columns
        if column.endswith(("_10k_kWh", "_MWh"))
        or column in {"deduction_ratio", "bridge_relative_error"}
    ]
    numeric = ledger[numeric_columns].apply(pd.to_numeric, errors="coerce")
    if numeric.isna().any().any() or (numeric < 0).any().any():
        raise AssertionError("官方三项排除电量台账包含空值、非数值或负值")
    if not ledger.identity_passed.astype(bool).all():
        raise AssertionError("官方三项排除电量台账存在恒等式未通过记录")
    return ledger


def _load_sample_mapping_entries(
    mapping_path: str | Path = SAMPLE_INDUSTRY_MAPPING_CSV,
) -> tuple[SampleIndustryMapping, ...]:
    path = Path(mapping_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    expected = {
        "sample_industry",
        "broad_sector",
        "is_traditional",
        "official_source_row",
        "official_label",
        "notes",
    }
    records: list[SampleIndustryMapping] = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if set(reader.fieldnames or ()) != expected:
            raise ValueError(f"sample mapping schema mismatch: {reader.fieldnames}")
        for line, raw in enumerate(reader, start=2):
            try:
                record = SampleIndustryMapping(
                    sample_industry=clean_text(raw["sample_industry"]),
                    broad_sector=clean_text(raw["broad_sector"]),
                    is_traditional=_parse_bool(raw["is_traditional"], field="is_traditional"),
                    official_source_row=int(clean_text(raw["official_source_row"])),
                    official_label=clean_text(raw["official_label"]),
                    notes=clean_text(raw["notes"]),
                )
            except (KeyError, TypeError, ValueError) as exc:
                raise ValueError(f"invalid sample mapping line {line}: {exc}") from exc
            if not record.sample_industry or not record.broad_sector:
                raise ValueError("sample industry and broad sector must be non-empty")
            records.append(record)
    industries = [row.sample_industry for row in records]
    if not records or len(industries) != len(set(industries)):
        raise ValueError("sample mapping must be non-empty and unique by sample_industry")
    traditional_count = sum(item.is_traditional for item in records)
    if traditional_count != TRADITIONAL_COMPONENT_COUNT:
        raise ValueError(
            f"sample mapping must contain exactly {TRADITIONAL_COMPONENT_COUNT} "
            f"traditional fine industries, got {traditional_count}"
        )
    return tuple(records)


def _load_official_mapping_entries(
    mapping_path: str | Path = OFFICIAL_2025_MAPPING_CSV,
) -> tuple[OfficialMappingEntry, ...]:
    path = Path(mapping_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    expected = {
        "source_row",
        "official_label",
        "classification",
        "accounting_role",
        "department",
        "split_parent_label",
        "representative_lookup_allowed",
        "notes",
    }
    records: list[OfficialMappingEntry] = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if set(reader.fieldnames or ()) != expected:
            raise ValueError(f"official mapping schema mismatch: {reader.fieldnames}")
        for line, raw in enumerate(reader, start=2):
            try:
                record = OfficialMappingEntry(
                    source_row=int(clean_text(raw["source_row"])),
                    official_label=clean_text(raw["official_label"]),
                    classification=clean_text(raw["classification"]),
                    accounting_role=AccountingRole(clean_text(raw["accounting_role"])),
                    department=clean_text(raw["department"]),
                    split_parent_label=clean_text(raw["split_parent_label"]),
                    representative_lookup_allowed=_parse_bool(
                        raw["representative_lookup_allowed"],
                        field="representative_lookup_allowed",
                    ),
                    notes=clean_text(raw["notes"]),
                )
            except (KeyError, TypeError, ValueError) as exc:
                raise ValueError(f"invalid official mapping line {line}: {exc}") from exc
            if not record.classification:
                raise ValueError(f"official row {record.source_row} has empty accounting class")
            if record.accounting_role is AccountingRole.SUMMARY:
                if record.classification != "汇总项":
                    raise ValueError("summary accounting rows must be classified 汇总项")
            elif record.classification == "汇总项":
                raise ValueError("non-summary accounting rows cannot be classified 汇总项")
            if record.accounting_role is AccountingRole.SPLIT_COMPONENT:
                if not record.split_parent_label:
                    raise ValueError("split_component requires split_parent_label")
            elif record.split_parent_label:
                raise ValueError("split_parent_label is only valid for split_component rows")
            if record.accounting_role in {
                AccountingRole.LEAF,
                AccountingRole.SPLIT_PARENT,
                AccountingRole.SPLIT_COMPONENT,
            } and not record.department:
                raise ValueError("accounted official rows require a department")
            records.append(record)

    source_rows = [row.source_row for row in records]
    labels = [row.official_label for row in records]
    if not records or len(source_rows) != len(set(source_rows)):
        raise ValueError("official mapping must be non-empty and unique by source_row")
    if len(labels) != len(set(labels)):
        raise ValueError("official mapping must be unique by normalized official_label")
    return tuple(records)


def _match_official_rows(
    official_rows: Sequence[Official2025Row],
    mapping_entries: Sequence[OfficialMappingEntry],
) -> dict[int, Official2025Row]:
    row_by_number = {row.source_row: row for row in official_rows}
    mapping_by_number = {row.source_row: row for row in mapping_entries}
    missing = sorted(set(row_by_number).difference(mapping_by_number))
    extra = sorted(set(mapping_by_number).difference(row_by_number))
    if missing or extra:
        raise ValueError(f"official mapping row coverage mismatch; missing={missing}, extra={extra}")
    mismatches = [
        (number, row_by_number[number].label, mapping_by_number[number].official_label)
        for number in sorted(row_by_number)
        if row_by_number[number].label != mapping_by_number[number].official_label
    ]
    if mismatches:
        raise ValueError(f"official mapping label mismatch: {mismatches[:10]}")
    return row_by_number


def _accounted_amounts(
    official_rows: Sequence[Official2025Row],
    mapping_entries: Sequence[OfficialMappingEntry],
) -> dict[int, float]:
    row_by_number = _match_official_rows(official_rows, mapping_entries)
    label_to_mapping = {entry.official_label: entry for entry in mapping_entries}
    label_to_row = {row.label: row for row in official_rows}
    components_by_parent: dict[str, list[OfficialMappingEntry]] = defaultdict(list)
    for entry in mapping_entries:
        if entry.accounting_role is AccountingRole.SPLIT_COMPONENT:
            components_by_parent[entry.split_parent_label].append(entry)

    amounts: dict[int, float] = {}
    for entry in mapping_entries:
        value = row_by_number[entry.source_row].value_2025
        if entry.accounting_role in {AccountingRole.SUMMARY, AccountingRole.DETAIL}:
            amounts[entry.source_row] = 0.0
        elif entry.accounting_role in {
            AccountingRole.LEAF,
            AccountingRole.SPLIT_COMPONENT,
        }:
            amounts[entry.source_row] = value
        elif entry.accounting_role is AccountingRole.SPLIT_PARENT:
            children = components_by_parent.get(entry.official_label, [])
            if not children:
                raise ValueError(f"split parent has no components: {entry.official_label}")
            component_total = sum(
                label_to_row[child.official_label].value_2025 for child in children
            )
            residual = value - component_total
            max_decimals = max(row.display_decimal_places for row in official_rows)
            tolerance = 0.5 * 10.0 ** (-max_decimals)
            if residual < -tolerance:
                raise ValueError(
                    f"split components exceed parent {entry.official_label}: "
                    f"{component_total} > {value}"
                )
            amounts[entry.source_row] = max(0.0, residual)
        else:  # pragma: no cover - Enum exhaustiveness guard
            raise AssertionError(entry.accounting_role)

    unknown_parents = sorted(set(components_by_parent).difference(label_to_mapping))
    if unknown_parents:
        raise ValueError(f"split components reference unknown parents: {unknown_parents}")
    invalid_parents = sorted(
        label
        for label in components_by_parent
        if label_to_mapping[label].accounting_role is not AccountingRole.SPLIT_PARENT
    )
    if invalid_parents:
        raise ValueError(f"split components reference non-split parents: {invalid_parents}")
    return amounts


def validate_official_balance(
    official_rows: Sequence[Official2025Row],
    mapping_entries: Sequence[OfficialMappingEntry],
    *,
    total_label: str = OFFICIAL_TOTAL_LABEL,
    department_total_labels: Mapping[str, str] = OFFICIAL_DEPARTMENT_TOTAL_LABELS,
    relative_tolerance: float = OFFICIAL_RELATIVE_TOLERANCE,
) -> OfficialBalanceResult:
    """Validate unique mapping and both department/full-society closure."""

    if not official_rows:
        raise ValueError("official_rows must not be empty")
    row_by_label = {row.label: row for row in official_rows}
    if total_label not in row_by_label:
        raise ValueError(f"official total label is missing: {total_label}")
    amounts = _accounted_amounts(official_rows, mapping_entries)
    max_decimals = max(row.display_decimal_places for row in official_rows)
    absolute_tolerance = 0.5 * 10.0 ** (-max_decimals)

    amount_by_classification: dict[str, float] = defaultdict(float)
    amount_by_department: dict[str, float] = defaultdict(float)
    for entry in mapping_entries:
        amount = amounts[entry.source_row]
        if amount == 0:
            continue
        amount_by_classification[entry.classification] += amount
        amount_by_department[entry.department] += amount

    department_balances: list[DepartmentBalance] = []
    for department, target in department_total_labels.items():
        if target not in row_by_label:
            raise ValueError(f"department target label is missing: {target}")
        target_value = row_by_label[target].value_2025
        mapped_value = float(amount_by_department.get(department, 0.0))
        absolute_error = abs(mapped_value - target_value)
        relative_error = absolute_error / target_value if target_value else absolute_error
        passed = absolute_error <= absolute_tolerance and relative_error <= relative_tolerance
        department_balances.append(
            DepartmentBalance(
                department=department,
                target_label=target,
                target_value_2025=target_value,
                mapped_value_2025=mapped_value,
                absolute_error=absolute_error,
                relative_error=relative_error,
                passed=passed,
            )
        )

    total_value = row_by_label[total_label].value_2025
    mapped_total = float(sum(amount_by_classification.values()))
    absolute_error = abs(mapped_total - total_value)
    relative_error = absolute_error / total_value if total_value else absolute_error
    total_passed = absolute_error <= absolute_tolerance and relative_error <= relative_tolerance
    passed = total_passed and all(item.passed for item in department_balances)
    result = OfficialBalanceResult(
        unit=official_rows[0].unit,
        total_label=total_label,
        total_value_2025=total_value,
        mapped_value_2025=mapped_total,
        absolute_tolerance=absolute_tolerance,
        relative_tolerance=relative_tolerance,
        absolute_error=absolute_error,
        relative_error=relative_error,
        passed=passed,
        amount_by_classification=dict(sorted(amount_by_classification.items())),
        amount_by_department=dict(sorted(amount_by_department.items())),
        department_balances=tuple(department_balances),
    )
    if not result.passed:
        failed_departments = [
            item.department for item in result.department_balances if not item.passed
        ]
        raise ValueError(
            "official 2025 mapping does not close: "
            f"total_abs_error={absolute_error:.12g}, "
            f"total_rel_error={relative_error:.12g}, "
            f"failed_departments={failed_departments}"
        )
    return result


def validate_sample_catalog(
    catalog: Sequence[SampleSheetRecord],
    mapping_entries: Sequence[SampleIndustryMapping],
) -> None:
    mapped = {row.sample_industry for row in mapping_entries}
    observed = {row.industry for row in catalog}
    missing = sorted(observed.difference(mapped))
    stale = sorted(mapped.difference(observed))
    if missing or stale:
        raise ValueError(f"sample industry mapping mismatch; missing={missing}, stale={stale}")
    broad_by_industry = {row.sample_industry: row.broad_sector for row in mapping_entries}
    conflicts = sorted(
        {
            row.industry
            for row in catalog
            if broad_by_industry[row.industry] != row.broad_sector
        }
    )
    if conflicts:
        raise ValueError(f"sample broad-sector mapping mismatch: {conflicts}")


def sample_industry_energy_records(
    sample_mapping: Sequence[SampleIndustryMapping],
    official_rows: Sequence[Official2025Row],
    official_mapping: Sequence[OfficialMappingEntry] | None = None,
) -> tuple[SampleIndustryEnergy, ...]:
    """Resolve each sample fine industry to one exact official 2025 cell."""

    row_by_number = {row.source_row: row for row in official_rows}
    mapping_by_number = (
        {entry.source_row: entry for entry in official_mapping}
        if official_mapping is not None
        else {}
    )
    accounted = (
        _accounted_amounts(official_rows, official_mapping)
        if official_mapping is not None
        else {}
    )
    output: list[SampleIndustryEnergy] = []
    for item in sample_mapping:
        source = row_by_number.get(item.official_source_row)
        if source is None:
            raise ValueError(
                f"sample industry {item.sample_industry} references missing official row "
                f"{item.official_source_row}"
            )
        if source.label != item.official_label:
            raise ValueError(
                f"sample industry {item.sample_industry} official label mismatch: "
                f"{source.label!r} != {item.official_label!r}"
            )
        official_entry = mapping_by_number.get(item.official_source_row)
        if (
            item.is_traditional
            and official_entry is not None
            and official_entry.accounting_role is AccountingRole.SPLIT_PARENT
        ):
            energy = accounted[item.official_source_row]
            energy_basis = "split_parent_residual"
        else:
            energy = source.value_2025
            energy_basis = "source_row_value"
        output.append(
            SampleIndustryEnergy(
                sample_industry=item.sample_industry,
                broad_sector=item.broad_sector,
                is_traditional=item.is_traditional,
                official_source_row=item.official_source_row,
                official_label=item.official_label,
                official_energy_2025=energy,
                official_energy_basis=energy_basis,
                unit=source.unit,
            )
        )
    return tuple(output)


def sample_industry_energy_frame(
    *,
    sample_mapping_path: str | Path = SAMPLE_INDUSTRY_MAPPING_CSV,
    official_mapping_path: str | Path = OFFICIAL_2025_MAPPING_CSV,
    official_workbook: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
) -> pd.DataFrame:
    """Return the exact table consumed by representative-source selection."""

    sample_mapping = _load_sample_mapping_entries(sample_mapping_path)
    official_rows = read_official_2025_rows(official_workbook)
    official_mapping = _load_official_mapping_entries(official_mapping_path)
    records = sample_industry_energy_records(
        sample_mapping, official_rows, official_mapping
    )
    return pd.DataFrame(
        [
            {
                "industry": row.sample_industry,
                "sample_industry": row.sample_industry,
                "broad_sector": row.broad_sector,
                "modeling_role": "传统细分行业" if row.is_traditional else "后续新型负荷模块样本",
                "is_traditional": row.is_traditional,
                "official_source_row": row.official_source_row,
                "official_label": row.official_label,
                "official_energy_2025": row.official_energy_2025,
                "official_energy_basis": row.official_energy_basis,
                "unit": row.unit,
            }
            for row in records
        ]
    )


def sample_industry_energy_year_frame(
    year: int,
    *,
    sample_mapping_path: str | Path = SAMPLE_INDUSTRY_MAPPING_CSV,
    official_mapping_path: str | Path = OFFICIAL_2025_MAPPING_CSV,
    official_workbook: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
) -> pd.DataFrame:
    """Resolve the 25 sample rows against another official historical year."""

    sample_mapping = _load_sample_mapping_entries(sample_mapping_path)
    official_rows = read_official_year_rows(official_workbook, year=int(year))
    official_mapping = _load_official_mapping_entries(official_mapping_path)
    records = sample_industry_energy_records(
        sample_mapping, official_rows, official_mapping
    )
    return pd.DataFrame(
        [
            {
                "year": int(year),
                "industry": row.sample_industry,
                "broad_sector": row.broad_sector,
                "modeling_role": (
                    "传统细分行业" if row.is_traditional else "后续新型负荷模块样本"
                ),
                "is_traditional": row.is_traditional,
                "official_source_row": row.official_source_row,
                "official_label": row.official_label,
                "official_energy": row.official_energy_2025,
                "official_energy_basis": row.official_energy_basis,
                "unit": row.unit,
            }
            for row in records
        ]
    )


def load_sample_mapping(
    mapping_path: str | Path = SAMPLE_INDUSTRY_MAPPING_CSV,
) -> pd.DataFrame:
    """Public sample-industry mapping interface."""

    records = _load_sample_mapping_entries(mapping_path)
    return pd.DataFrame(
        [
            {
                "industry": row.sample_industry,
                "sample_industry": row.sample_industry,
                "broad_sector": row.broad_sector,
                "modeling_role": "传统细分行业" if row.is_traditional else "后续新型负荷模块样本",
                "is_traditional": row.is_traditional,
                "official_source_row": row.official_source_row,
                "official_label": row.official_label,
                "notes": row.notes,
            }
            for row in records
        ]
    )


def load_official_mapping(
    mapping_path: str | Path = OFFICIAL_2025_MAPPING_CSV,
    official_workbook: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
) -> pd.DataFrame:
    """Return source-row accounting evidence without creating model groups.

    The source workbook's historical classification column is reduced to four
    audit-only accounting classes.  It is never used to allocate electricity or
    to choose shape parameters.
    """

    mapping_entries = _load_official_mapping_entries(mapping_path)
    official_rows = read_official_2025_rows(official_workbook)
    validate_official_balance(official_rows, mapping_entries)
    row_by_number = _match_official_rows(official_rows, mapping_entries)
    accounted = _accounted_amounts(official_rows, mapping_entries)
    frame = pd.DataFrame(
        [
            {
                "source_row": entry.source_row,
                "source_label": entry.official_label,
                "department": entry.department,
                "accounting_class": (
                    entry.classification
                    if entry.classification in {"汇总项", "新型负荷", "排除项"}
                    else "传统及其他计量项"
                ),
                "accounting_role": entry.accounting_role.value,
                "electricity_2025": row_by_number[entry.source_row].value_2025,
                "accounted_electricity_2025": accounted[entry.source_row],
                "unit": row_by_number[entry.source_row].unit,
                "is_leaf": entry.accounting_role
                in {AccountingRole.LEAF, AccountingRole.SPLIT_COMPONENT},
                "is_accounted": entry.accounting_role
                in {
                    AccountingRole.LEAF,
                    AccountingRole.SPLIT_PARENT,
                    AccountingRole.SPLIT_COMPONENT,
                },
                "split_parent_label": entry.split_parent_label,
                "representative_lookup_allowed": entry.representative_lookup_allowed,
                "notes": entry.notes,
            }
            for entry in mapping_entries
        ]
    )
    return frame.sort_values("source_row", kind="mergesort").reset_index(drop=True)


def official_industry_energy(mapping: pd.DataFrame) -> dict[str, float]:
    """Return a unique official-energy lookup from an enriched mapping frame.

    For an official mapping frame the keys are source labels.  For a sample
    energy frame the keys are sample industries.  Duplicate keys are rejected.
    """

    if {"industry", "official_energy_2025"}.issubset(mapping.columns):
        key, value = "industry", "official_energy_2025"
    elif {"source_label", "electricity_2025"}.issubset(mapping.columns):
        key, value = "source_label", "electricity_2025"
    else:
        raise ValueError(
            "mapping must contain either industry/official_energy_2025 or "
            "source_label/electricity_2025"
        )
    data = mapping[[key, value]].copy()
    data[key] = data[key].map(clean_text)
    if data[key].eq("").any() or data[key].duplicated().any():
        raise ValueError(f"official industry energy keys must be non-empty and unique: {key}")
    data[value] = pd.to_numeric(data[value], errors="coerce")
    if (~data[value].map(math.isfinite)).any() or (data[value] < 0).any():
        raise ValueError("official industry electricity must be finite and non-negative")
    return dict(zip(data[key], data[value].astype(float)))


def load_and_validate_inputs(
    *,
    load_workbook: str | Path = DEFAULT_LOAD_WORKBOOK,
    official_workbook: str | Path = DEFAULT_OFFICIAL_WORKBOOK,
    sample_mapping_path: str | Path = SAMPLE_INDUSTRY_MAPPING_CSV,
    official_mapping_path: str | Path = OFFICIAL_2025_MAPPING_CSV,
) -> dict[str, object]:
    """One-stop formal input gate used before any model calculation."""

    manifest = build_input_manifest(
        {
            "load_workbook": load_workbook,
            "official_2025_workbook": official_workbook,
            "sample_industry_mapping": sample_mapping_path,
            "official_row_mapping": official_mapping_path,
        }
    )
    catalog = read_sample_sheet_catalog(load_workbook)
    sample_entries = _load_sample_mapping_entries(sample_mapping_path)
    validate_sample_catalog(catalog, sample_entries)
    official_rows = read_official_2025_rows(official_workbook)
    official_entries = _load_official_mapping_entries(official_mapping_path)
    balance = validate_official_balance(official_rows, official_entries)
    excluded_energy_ledger = build_excluded_energy_ledger(official_workbook)
    energy_frame = sample_industry_energy_frame(
        sample_mapping_path=sample_mapping_path,
        official_mapping_path=official_mapping_path,
        official_workbook=official_workbook,
    )
    official_frame = load_official_mapping(official_mapping_path, official_workbook)
    return {
        "manifest": manifest,
        "sample_catalog": catalog,
        "sample_mapping": load_sample_mapping(sample_mapping_path),
        "official_mapping": official_frame,
        "official_balance": balance,
        "excluded_energy_ledger": excluded_energy_ledger,
        "sample_industry_energy": energy_frame,
    }


__all__ = [
    "build_excluded_energy_ledger",
    "build_input_manifest",
    "clean_text",
    "load_and_validate_inputs",
    "load_official_mapping",
    "load_sample_mapping",
    "official_industry_energy",
    "parse_sample_sheet_name",
    "read_official_2025_rows",
    "read_official_year_rows",
    "read_sample_sheet_catalog",
    "sample_industry_energy_frame",
    "sample_industry_energy_year_frame",
    "sample_industry_energy_records",
    "sha256_file",
    "validate_official_balance",
    "validate_sample_catalog",
    "verify_input_manifest",
]
